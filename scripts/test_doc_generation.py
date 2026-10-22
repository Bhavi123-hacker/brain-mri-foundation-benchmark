import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

doc = docx.Document()
doc.add_heading("SHIFT-MRI V4 Test", 0)
doc.add_paragraph("Testing python-docx generation.")
doc.save("test_docx.docx")
print("DOCX generated successfully.")

styles = getSampleStyleSheet()
pdf = SimpleDocTemplate("test_pdf.pdf", pagesize=letter)
story = [Paragraph("SHIFT-MRI V4 PDF Test", styles["Heading1"]), Spacer(1, 10), Paragraph("Testing reportlab generation.", styles["Normal"])]
pdf.build(story)
print("PDF generated successfully.")
